#!/usr/bin/env python3
"""Offline regression checks for distribution auditing."""
import importlib.util
import tempfile
import unittest
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

spec=importlib.util.spec_from_file_location('distribution',Path(__file__).with_name('distribution_check.py'))
distribution=importlib.util.module_from_spec(spec); spec.loader.exec_module(distribution)


class DistributionTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory(); self.path=Path(self.tmp.name)/'test.docx'
        self.doc=Document()
        for field in ('author','last_modified_by','comments','keywords'):
            setattr(self.doc.core_properties,field,'')
        self.doc.add_paragraph('Neutral technical content.')
    def tearDown(self):
        self.tmp.cleanup()
    def audit(self,terms=()):
        self.doc.save(self.path)
        return distribution.audit(self.path,terms)
    def test_neutral_docx_passes(self):
        self.assertTrue(self.audit()['ok'])
    def test_author_metadata_is_rejected(self):
        self.doc.core_properties.author='Private Example Name'
        self.assertFalse(self.audit()['ok'])
    def test_identity_split_across_runs_is_rejected(self):
        p=self.doc.add_paragraph(); p.add_run('PRIVATE_'); p.add_run('MARKER')
        self.assertFalse(self.audit(['PRIVATE_MARKER'])['ok'])
    def test_header_identity_is_rejected(self):
        self.doc.sections[0].header.paragraphs[0].text='PRIVATE_MARKER'
        self.assertFalse(self.audit(['PRIVATE_MARKER'])['ok'])
    def test_hidden_text_is_rejected(self):
        self.doc.paragraphs[0].runs[0].font.hidden=True
        self.assertFalse(self.audit()['ok'])
    def test_revision_is_rejected(self):
        node=OxmlElement('w:ins'); node.set(qn('w:id'),'1')
        self.doc.paragraphs[0]._p.append(node)
        self.assertFalse(self.audit()['ok'])


if __name__=='__main__': unittest.main(verbosity=2)

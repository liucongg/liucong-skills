#!/usr/bin/env python3
"""Build and inspect editable illustrated DOCX books; never certify science or taste."""
from __future__ import annotations
import argparse
import base64
import hashlib
import importlib.util
import json
import math
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

MAX_SVG = 12 * 1024 * 1024
ALLOWED_KINDS = {"official", "textbook", "paper", "standard", "institutional", "reference", "community"}


def dump(data, out=None):
    text = json.dumps(data, ensure_ascii=False, indent=2)
    if out:
        p = Path(out)
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_text(text + "\n", encoding="utf-8")
    print(text)
    return data


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def confined(root, value):
    if not isinstance(value, str) or not value:
        raise ValueError("asset path must be a nonempty string")
    p = (root / value).resolve()
    if not p.is_relative_to(root.resolve()):
        raise ValueError(f"asset path escapes project: {value}")
    return p


def check_svg(path):
    path = Path(path)
    raw = path.read_bytes()
    if len(raw) > MAX_SVG:
        raise ValueError("SVG exceeds size limit")
    low = raw.lower()
    if b"<!doctype" in low or b"<!entity" in low:
        raise ValueError("SVG DTD/entity forbidden")
    root = ET.fromstring(raw)
    if root.tag.rsplit("}", 1)[-1] != "svg":
        raise ValueError("not an SVG document")
    if not root.get("viewBox"):
        raise ValueError("SVG needs viewBox")
    box = [float(x) for x in re.split(r"[\s,]+", root.get("viewBox").strip())]
    if len(box) != 4 or box[2] <= 0 or box[3] <= 0 or not all(math.isfinite(x) for x in box):
        raise ValueError("invalid SVG viewBox")
    for el in root.iter():
        local = el.tag.rsplit("}", 1)[-1].lower()
        if local in {"script", "foreignobject", "animate", "animatetransform", "set"}:
            raise ValueError(f"non-static SVG element: {local}")
        for key, val in el.attrib.items():
            key = key.rsplit("}", 1)[-1].lower()
            if key.startswith("on"):
                raise ValueError("SVG event attribute forbidden")
            if key in {"href", "src"} and not (val.startswith("#") or val.startswith("data:image/")):
                raise ValueError("SVG external resource forbidden")
        text = " ".join(el.attrib.values()) + " " + (el.text or "")
        if re.search(r"@import|javascript\s*:|url\(\s*['\"]?(?!#|data:image/)", text, re.I):
            raise ValueError("SVG external CSS/resource forbidden")
    return root, box


def chrome_path():
    names = ["google-chrome", "google-chrome-stable", "chromium", "chromium-browser", "chrome"]
    for name in names:
        if shutil.which(name):
            return shutil.which(name)
    candidates = [
        Path("/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"),
        Path("/Applications/Chromium.app/Contents/MacOS/Chromium"),
    ]
    for key in ("PROGRAMFILES", "PROGRAMFILES(X86)", "LOCALAPPDATA"):
        if os.environ.get(key):
            candidates.append(Path(os.environ[key]) / "Google/Chrome/Application/chrome.exe")
    return next((str(p) for p in candidates if p.is_file()), None)


def soffice_path():
    return shutil.which("soffice") or shutil.which("libreoffice") or next((str(p) for p in [
        Path("/Applications/LibreOffice.app/Contents/MacOS/soffice"),
        Path(os.environ.get("PROGRAMFILES", ".")) / "LibreOffice/program/soffice.exe",
    ] if p.is_file()), None)


def preflight():
    return dump({"python": sys.executable, "node": shutil.which("node"),
        "modules": {m: bool(importlib.util.find_spec(m)) for m in ("docx", "PIL", "fitz", "cairosvg")},
        "rsvg_convert": shutil.which("rsvg-convert"), "chrome": chrome_path(),
        "soffice": soffice_path(), "installs_performed": False})


def indexed(items, label, errors):
    out = {}
    for item in items:
        ident = item.get("id")
        if not isinstance(ident, str) or not ident or ident in out:
            errors.append(f"missing/duplicate {label} id: {ident}")
        else:
            out[ident] = item
    return out


def validate_book(manifest, require_raster=True):
    path = Path(manifest).resolve()
    data = json.loads(path.read_text(encoding="utf-8"))
    root, errors = path.parent, []
    for field in ("title", "reader", "mode", "scope"):
        if not data.get(field):
            errors.append(f"missing book field: {field}")
    if data.get("mode") not in {"real-structure", "concept-analogy"}:
        errors.append("mode must be real-structure or concept-analogy")
    sources = indexed(data.get("sources", []), "source", errors)
    claims = indexed(data.get("claims", []), "claim", errors)
    parts = indexed(data.get("parts", []), "part", errors)
    figures = indexed(data.get("figures", []), "figure", errors)
    verified = [s for s in sources.values() if s.get("verified") is True and s.get("excerpt")]
    required_sources = data.get("min_verified_sources", 1)
    if not isinstance(required_sources, int) or isinstance(required_sources, bool) or required_sources < 1:
        errors.append("min_verified_sources must be a positive integer")
        required_sources = 1
    if len(verified) < required_sources:
        errors.append(f"need at least {required_sources} actually-read sources with excerpts")
    if not any(s.get("kind") in {"official", "textbook", "paper", "standard", "institutional"} for s in verified):
        errors.append("need a direct authoritative source")
    for sid, s in sources.items():
        if not re.match(r"https?://", s.get("url", "")) or not s.get("title"):
            errors.append(f"{sid}: source title/URL missing")
        if s.get("kind") not in ALLOWED_KINDS:
            errors.append(f"{sid}: unknown source kind")
        if s.get("verified") and (not s.get("excerpt") or not s.get("retrieved_on")):
            errors.append(f"{sid}: verified source lacks excerpt/retrieved_on")
    def refs(item, name, table, owner):
        value = item.get(name, [])
        if not isinstance(value, list):
            errors.append(f"{owner}: {name} must be a list")
            return []
        for ident in value:
            if ident not in table:
                errors.append(f"{owner}: unknown {name} {ident}")
        return value
    for cid, claim in claims.items():
        ids = refs(claim, "source_ids", sources, cid)
        status = claim.get("status")
        if not claim.get("text") or status not in {"verified", "assumption", "unverified"}:
            errors.append(f"{cid}: invalid claim")
        if claim.get("critical", True) and status != "verified":
            if not (status == "assumption" and claim.get("test_evidence") and claim.get("assumption_notice")):
                errors.append(f"{cid}: critical claim not verified or tested explicit assumption")
        if status == "verified" and (not ids or any(not sources.get(i, {}).get("verified") for i in ids)):
            errors.append(f"{cid}: verified claim needs verified sources")
    for fid, fig in figures.items():
        method = fig.get("method", "svg")
        role = fig.get("role", "technical")
        if method not in {"svg", "existing", "generated", "hybrid"} or role not in {"technical", "context"}:
            errors.append(f"{fid}: invalid media method/role")
        required = ["question", "kind", "caption", "alt_text", "png"]
        if method == "svg":
            required.append("svg")
        for field in required:
            if not fig.get(field):
                errors.append(f"{fid}: missing {field}")
        for name, table in (("part_ids", parts), ("claim_ids", claims), ("source_ids", sources)):
            refs(fig, name, table, fid)
        if role == "technical" and not fig.get("claim_ids"):
            errors.append(f"{fid}: no explanation claim")
        if fig.get("provenance") != "original" and not fig.get("license"):
            errors.append(f"{fid}: external figure needs license")
        try:
            if fig.get("svg"):
                check_svg(confined(root, fig["svg"]))
            png = confined(root, fig.get("png"))
            if require_raster:
                from PIL import Image
                with Image.open(png) as im:
                    if im.format != "PNG":
                        errors.append(f"{fid}: raster must be PNG")
                    width_mm = float(fig.get("width_mm", 164))
                    if not 20 <= width_mm <= 174:
                        errors.append(f"{fid}: width_mm outside portrait content width")
                    if im.width / (width_mm / 25.4) < (250 if role == "technical" else 150):
                        errors.append(f"{fid}: insufficient raster pixels for print size")
        except (OSError, ValueError, TypeError, ET.ParseError) as exc:
            errors.append(f"{fid}: {exc}")
    if data.get("schema_version") == "0.3" or data.get("media_manifest"):
        try:
            media_path = confined(root, data.get("media_manifest"))
            spec = importlib.util.spec_from_file_location("book_media", Path(__file__).with_name("media.py"))
            module = importlib.util.module_from_spec(spec); spec.loader.exec_module(module)
            audit = module.audit(media_path)
            errors.extend("media: " + e for e in audit["errors"])
            registry = json.loads(media_path.read_text(encoding="utf-8"))
            assets = {a["id"]: a for a in registry.get("assets", [])}
            for fid, fig in figures.items():
                asset = assets.get(fig.get("asset_id"))
                if not asset:
                    errors.append(f"{fid}: no linked media asset")
                elif (confined(media_path.parent, asset["path"]) != confined(root, fig["png"]) or
                      asset.get("method") != fig.get("method", "svg") or
                      asset.get("role") != fig.get("role", "technical")):
                    errors.append(f"{fid}: media path/method/role mismatch")
        except (OSError, ValueError, TypeError, KeyError) as exc:
            errors.append(f"media manifest: {exc}")
    used_figures, text = set(), []
    pages = data.get("pages", [])
    delivery_kind = data.get("delivery_kind", "production")
    if delivery_kind not in {"production", "sample-chapter", "technical-smoke"}:
        errors.append("invalid delivery_kind")
    if delivery_kind == "production":
        if len(pages) < 6:
            errors.append("production requires at least 6 substantive editorial units")
        if len(figures) < 6:
            errors.append("production requires at least 6 original explanatory plates")
        if len({f.get("kind") for f in figures.values()}) < 4:
            errors.append("production requires at least 4 distinct explanation view types")
    if delivery_kind == "sample-chapter":
        if not data.get("sample_objectives") or not data.get("sample_limits"):
            errors.append("sample-chapter needs explicit objectives and limits")
        if not any(f.get("role", "technical") == "technical" for f in figures.values()):
            errors.append("sample-chapter needs an explanatory plate, not only decoration")
    if not pages:
        errors.append("book has no content pages")
    for page in pages:
        if not page.get("title") or not page.get("blocks"):
            errors.append("page requires title and nonempty blocks")
        text.extend([page.get("title", ""), page.get("lead", "")])
        for block in page.get("blocks", []):
            typ = block.get("type")
            if typ == "figure":
                ident = block.get("id")
                if ident not in figures:
                    errors.append(f"unknown figure block: {ident}")
                used_figures.add(ident)
            elif typ in {"paragraph", "callout"}:
                if not block.get("text"):
                    errors.append(f"empty {typ} block")
                text.append(block.get("text", ""))
            elif typ == "table":
                headers = block.get("headers", [])
                if not headers or any(len(r) != len(headers) for r in block.get("rows", [])):
                    errors.append("table columns mismatch")
                text.append(json.dumps(block, ensure_ascii=False))
            elif typ == "steps":
                if not block.get("items"):
                    errors.append("steps block needs items")
                text.extend(str(x) for x in block.get("items", []))
            else:
                errors.append(f"unsupported block type: {typ}")
    for sid in re.findall(r"\[(S\w+)\]", " ".join(text)):
        if sid not in sources:
            errors.append(f"unknown text source marker {sid}")
    if set(figures) - used_figures:
        errors.append("unused figure definitions: " + ",".join(sorted(set(figures) - used_figures)))
    if delivery_kind == "production" and len("".join(text)) < 1000:
        errors.append("production editable manuscript is too thin")
    report = {"ok": not errors, "errors": errors, "planned_pages": len(pages),
              "delivery_kind": delivery_kind, "manuscript_characters": len("".join(text)),
              "figures": len(figures), "verified_sources": len(verified),
              "scope": "structural/evidence-declaration checks, not scientific or visual approval"}
    return data, report


def raster(svg, out, width=2000):
    from PIL import Image
    _, box = check_svg(svg)
    width = int(width)
    height = round(width * box[3] / box[2])
    if width < 1 or height < 1 or width > 6000 or height > 6000:
        raise ValueError("raster dimensions outside limits")
    out = Path(out).resolve()
    out.parent.mkdir(parents=True, exist_ok=True)
    rsvg = shutil.which("rsvg-convert")
    backend = ""
    if rsvg:
        subprocess.run([rsvg, "-w", str(width), "-h", str(height), "-o", str(out), str(Path(svg).resolve())], check=True, timeout=90)
        backend = "rsvg-convert"
    elif importlib.util.find_spec("cairosvg"):
        import cairosvg
        cairosvg.svg2png(url=str(Path(svg).resolve()), write_to=str(out), output_width=width, output_height=height)
        backend = "cairosvg"
    else:
        chrome = chrome_path()
        if not chrome:
            raise RuntimeError("no SVG raster backend (rsvg-convert, CairoSVG, or Chrome)")
        with tempfile.TemporaryDirectory(prefix="manual-svg-") as tmp:
            tmp = Path(tmp)
            payload = base64.b64encode(Path(svg).read_bytes()).decode("ascii")
            html = '<!doctype html><meta charset="utf-8"><style>html,body{margin:0;width:100%;height:100%;background:white;overflow:hidden}img{width:100%;height:100%;object-fit:contain;display:block}</style><img alt="" src="data:image/svg+xml;base64,' + payload + '">'
            page = tmp / "image.html"
            page.write_text(html, encoding="utf-8")
            cmd = [chrome, "--headless=new", "--no-first-run", "--no-default-browser-check", "--hide-scrollbars",
                "--force-device-scale-factor=1", f"--user-data-dir={tmp / 'profile'}", f"--window-size={width},{height}",
                "--virtual-time-budget=1200", f"--screenshot={out}", page.as_uri()]
            subprocess.run(cmd, check=True, capture_output=True, timeout=90)
        backend = "Chrome"
    with Image.open(out) as image:
        actual = image.size
        if actual != (width, height):
            raise RuntimeError(f"raster size mismatch: {actual} vs {(width, height)}")
    return dump({"ok": True, "backend": backend, "out": str(out), "pixels": actual, "svg_sha256": digest(svg)})


def field(paragraph, code):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = code
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for el in (begin, instr, end):
        run._r.append(el)


def build(manifest, out):
    from docx import Document
    from docx.shared import Mm, Pt, RGBColor
    from docx.oxml.ns import qn
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    book, report = validate_book(manifest)
    if not report["ok"]:
        dump(report)
        raise ValueError("book validation failed")
    base = Path(manifest).resolve().parent
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.top_margin, sec.bottom_margin = Mm(18), Mm(18)
    sec.left_margin, sec.right_margin = Mm(18), Mm(18)
    font = book.get("font", "PingFang SC" if sys.platform == "darwin" else "Noto Sans CJK SC")
    design = book.get("design", {})
    accent = design.get("accent", "39766F")
    ink = design.get("ink", "263338")
    for name, size, color in [("Normal", design.get("body_pt", 10.5), ink), ("Heading 1", design.get("heading_pt", 21), accent),
        ("Heading 2", 14, accent), ("Caption", 8.5, "58685F"), ("Title", 29, ink)]:
        style = doc.styles[name]
        style.font.name, style.font.size = font, Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), font)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = design.get("line_spacing", 1.32)
    doc.styles["Heading 1"].paragraph_format.keep_with_next = True
    header = sec.header.paragraphs[0]
    header.text = book.get("series", "图解原理手册") + " · " + book["title"]
    header.runs[0].font.size = Pt(8)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    field(footer, "PAGE")
    footer.add_run(" / ")
    field(footer, "NUMPAGES")
    doc.core_properties.title = book["title"]
    doc.core_properties.subject = book["scope"]
    # Public documents do not inherit identity from legacy manifests or templates.
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.comments = ""
    doc.core_properties.keywords = ""
    if book.get("include_toc"):
        doc.add_heading(book["title"], 0)
        field(doc.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u')
        doc.add_page_break()
    figures = {f["id"]: f for f in book["figures"]}
    for index, page in enumerate(book["pages"]):
        break_this = page.get("break_before", index > 0 and book.get("page_layout") == "fixed")
        if page.get("kicker"):
            p = doc.add_paragraph(page["kicker"])
            p.paragraph_format.page_break_before = bool(break_this)
            p.runs[0].font.size = Pt(9)
            p.runs[0].font.color.rgb = RGBColor.from_string("527D83")
            p.paragraph_format.keep_with_next = True
        heading = doc.add_heading(page["title"], 1)
        if not page.get("kicker"):
            heading.paragraph_format.page_break_before = bool(break_this)
        if page.get("lead"):
            p = doc.add_paragraph(page["lead"])
            p.runs[0].bold = True
        for block in page["blocks"]:
            typ = block["type"]
            if typ == "paragraph":
                doc.add_paragraph(block["text"])
            elif typ == "callout":
                p = doc.add_paragraph()
                p.add_run(block.get("label", "读图提示") + "  ").bold = True
                p.add_run(block["text"])
            elif typ == "steps":
                for number, text in enumerate(block["items"], 1):
                    p = doc.add_paragraph()
                    p.add_run(f"{number:02d}  ").bold = True
                    p.add_run(str(text))
            elif typ == "figure":
                f = figures[block["id"]]
                p = doc.add_paragraph()
                p.paragraph_format.keep_with_next = True
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                pic = p.add_run().add_picture(str(confined(base, f["png"])), width=Mm(float(f.get("width_mm", 164))))
                pic._inline.docPr.set("descr", f["alt_text"])
                pic._inline.docPr.set("name", f["id"])
                caption = doc.add_paragraph(f["id"] + "  " + f["caption"], "Caption")
                caption.paragraph_format.keep_together = True
            elif typ == "table":
                table = doc.add_table(rows=1, cols=len(block["headers"]))
                table.style = "Table Grid"
                table.autofit = False
                col_width = Mm(174 / len(block["headers"]))
                for col in table.columns:
                    col.width = col_width
                for cell, title in zip(table.rows[0].cells, block["headers"]):
                    cell.width = col_width
                    cell.text = str(title)
                    for run in cell.paragraphs[0].runs:
                        run.bold = True
                for row in block.get("rows", []):
                    cells = table.add_row().cells
                    for cell, value in zip(cells, row):
                        cell.width, cell.text = col_width, str(value)
                from docx.oxml import OxmlElement
                for ri, row in enumerate(table.rows):
                    no_split = OxmlElement("w:cantSplit")
                    row._tr.get_or_add_trPr().append(no_split)
                    for cell in row.cells:
                        if ri == 0:
                            shading = OxmlElement("w:shd")
                            shading.set(qn("w:fill"), "E2EBE4")
                            cell._tc.get_or_add_tcPr().append(shading)
                        for paragraph in cell.paragraphs:
                            paragraph.paragraph_format.space_after = Pt(3)
                            paragraph.paragraph_format.space_before = Pt(2)
                            for run in paragraph.runs:
                                run.font.size = Pt(9)
    if book.get("include_sources", True):
        source_heading = doc.add_heading("参考资料与图版说明", 2)
        source_heading.paragraph_format.page_break_before = bool(book.get("sources_on_new_page", False))
        for source in book["sources"]:
            if source.get("verified"):
                p = doc.add_paragraph(f"[{source['id']}] {source['title']}\n{source['url']}")
                for run in p.runs:
                    run.font.size = Pt(8)
        doc.add_paragraph(book.get("figure_notice", "图版为原理示意，不用于制造、维修或尺寸测量。正文可编辑；插图为由保留的SVG原稿生成的PNG兼容图。"), "Caption")
    out = Path(out).resolve()
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    return dump({"ok": True, "out": str(out), "sha256": digest(out), "planned_pages": len(book["pages"]),
                 "embedded_format": "PNG", "svg_sources_retained": any(f.get("svg") for f in book["figures"]), "layout_reviewed": False})


def inspect_docx(path, out=None):
    from docx import Document
    path = Path(path)
    missing, images, external = [], [], []
    with zipfile.ZipFile(path) as z:
        names = set(z.namelist())
        if sum(i.file_size for i in z.infolist()) > 150 * 1024 * 1024:
            raise ValueError("DOCX decompressed size limit")
        for name in names:
            if name.startswith("word/media/"):
                images.append(name)
            if name.endswith(".rels"):
                parent = Path(name).parent.parent
                for rel in ET.fromstring(z.read(name)):
                    target = rel.get("Target", "")
                    if rel.get("TargetMode") == "External":
                        external.append(target)
                        continue
                    target_path = os.path.normpath(str(parent / target)).replace("\\", "/").lstrip("/")
                    if target_path not in names:
                        missing.append({"part": name, "target": target})
    doc = Document(path)
    body = [p.text for p in doc.paragraphs if p.text]
    table_text = [c.text for t in doc.tables for row in t.rows for c in row.cells]
    chars = len("".join(body + table_text))
    report = {"ok": not missing and chars > 100 and bool(images), "sha256": digest(path),
        "paragraphs": len(doc.paragraphs), "body_characters": chars, "images": sorted(images),
        "tables": len(doc.tables), "missing_relationships": missing, "external_relationships": external,
        "scope": "package and editable text checks only"}
    return dump(report, out)


def render_docx(path, out_dir):
    import fitz
    from PIL import Image, ImageOps, ImageDraw
    path, out_dir = Path(path).resolve(), Path(out_dir).resolve()
    out_dir.mkdir(parents=True, exist_ok=True)
    exe = soffice_path()
    if not exe:
        raise RuntimeError("LibreOffice unavailable; DOCX layout NOT verified")
    with tempfile.TemporaryDirectory(prefix="manual-office-") as tmp:
        cmd = [exe, f"-env:UserInstallation={Path(tmp).as_uri()}", "--headless", "--convert-to", "pdf", "--outdir", str(out_dir), str(path)]
        completed = subprocess.run(cmd, check=True, capture_output=True, text=True, timeout=180)
    pdf = out_dir / (path.stem + ".pdf")
    if not pdf.exists() or pdf.stat().st_mtime + 2 < path.stat().st_mtime:
        raise RuntimeError("renderer did not produce a fresh PDF: " + completed.stdout)
    pages, thumbnails = [], []
    with fitz.open(pdf) as document:
        for index, page in enumerate(document):
            target = out_dir / f"page-{index + 1:03d}.png"
            pix = page.get_pixmap(matrix=fitz.Matrix(1.6, 1.6), alpha=False)
            pix.save(target)
            body_text = "".join(b[4] for b in page.get_text("blocks") if b[6] == 0 and b[1] >= 40 and b[3] <= page.rect.height - 40).strip()
            pages.append({"page": index + 1, "path": target.name, "text_chars": len(body_text), "reviewed": False})
            with Image.open(target) as im:
                thumb = ImageOps.contain(im.convert("RGB"), (300, 425))
            tile = Image.new("RGB", (320, 455), "#eeeeee")
            tile.paste(thumb, ((320 - thumb.width) // 2, 10))
            ImageDraw.Draw(tile).text((12, 438), str(index + 1), fill="black")
            thumbnails.append(tile)
    cols = min(3, max(1, len(thumbnails)))
    rows = math.ceil(len(thumbnails) / cols)
    sheet = Image.new("RGB", (cols * 320, rows * 455), "white")
    for index, tile in enumerate(thumbnails):
        sheet.paste(tile, ((index % cols) * 320, (index // cols) * 455))
    sheet.save(out_dir / "contact-sheet.png")
    report = {"ok": True, "renderer": "LibreOffice", "docx_sha256": digest(path), "pdf_sha256": digest(pdf),
        "page_count": len(pages), "pages": pages, "blank_text_pages": [p["page"] for p in pages if p["text_chars"] < 15],
        "layout_reviewed": False, "note": "Open final page images; do not mark reviewed from tool success alone."}
    dump(report, out_dir / "render.json")
    review = out_dir / "review.json"
    if not review.exists():
        review.write_text(json.dumps({"docx_sha256": digest(path), "pages": [{"page": p["page"], "reviewed": False, "issues": []} for p in pages],
            "science_reviewed": False, "science_notes": "", "limitations": []}, ensure_ascii=False, indent=2), encoding="utf-8")
    return report


def qa_book(manifest, docx, qa_dir):
    _, validation = validate_book(manifest)
    package = inspect_docx(docx)
    qa_dir = Path(qa_dir)
    errors = list(validation["errors"])
    if not package["ok"]:
        errors.append("DOCX package/text/image checks failed")
    current = digest(docx)
    try:
        render = json.loads((qa_dir / "render.json").read_text(encoding="utf-8"))
        review = json.loads((qa_dir / "review.json").read_text(encoding="utf-8"))
        if render.get("docx_sha256") != current or review.get("docx_sha256") != current:
            errors.append("QA does not match final DOCX hash")
        expected = set(range(1, render["page_count"] + 1))
        records = review.get("pages", [])
        if {p.get("page") for p in records} != expected or len(records) != len(expected):
            errors.append("review does not cover every final page exactly once")
        if any(p.get("reviewed") is not True or p.get("issues") for p in records):
            errors.append("pages not reviewed or unresolved issues")
        for p in render.get("pages", []):
            if not (qa_dir / p["path"]).is_file():
                errors.append(f"missing page preview: {p['path']}")
        if not review.get("science_reviewed") or not review.get("science_notes"):
            errors.append("scientific/meaning review not recorded")
        if render.get("blank_text_pages"):
            errors.append("blank/near-blank text pages need explicit resolution")
    except (OSError, ValueError, KeyError, TypeError) as exc:
        errors.append(f"missing/invalid QA evidence: {exc}")
    return dump({"ok": not errors, "errors": errors, "docx_sha256": current,
        "scope": "Evidence declarations and matching artifacts, not automatic scientific certification"}, qa_dir / "qa.json")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)
    sub.add_parser("preflight")
    p = sub.add_parser("validate"); p.add_argument("manifest"); p.add_argument("--without-raster", action="store_true")
    p = sub.add_parser("raster"); p.add_argument("svg"); p.add_argument("--out", required=True); p.add_argument("--width", type=int, default=2000)
    p = sub.add_parser("build"); p.add_argument("manifest"); p.add_argument("--out", required=True)
    p = sub.add_parser("render"); p.add_argument("docx"); p.add_argument("--out-dir", required=True)
    p = sub.add_parser("inspect"); p.add_argument("docx"); p.add_argument("--out")
    p = sub.add_parser("qa"); p.add_argument("manifest"); p.add_argument("docx"); p.add_argument("--qa-dir", required=True)
    args = parser.parse_args()
    try:
        if args.command == "preflight": result = preflight()
        elif args.command == "validate": _, result = validate_book(args.manifest, not args.without_raster); dump(result)
        elif args.command == "raster": result = raster(args.svg, args.out, args.width)
        elif args.command == "build": result = build(args.manifest, args.out)
        elif args.command == "render": result = render_docx(args.docx, args.out_dir)
        elif args.command == "inspect": result = inspect_docx(args.docx, args.out)
        else: result = qa_book(args.manifest, args.docx, args.qa_dir)
        return 0 if result.get("ok", True) else 1
    except Exception as exc:
        dump({"ok": False, "error": str(exc)})
        return 1


if __name__ == "__main__":
    raise SystemExit(main())

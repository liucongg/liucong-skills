#!/usr/bin/env python3
"""Local regression tests; deliberately do not fabricate external evidence."""
import copy
import importlib.util
import json
import tempfile
import unittest
from pathlib import Path

HERE = Path(__file__).resolve().parent
spec = importlib.util.spec_from_file_location("manual", HERE / "manual.py")
manual = importlib.util.module_from_spec(spec)
spec.loader.exec_module(manual)

class ManualTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.svg = self.root / "figure.svg"
        self.svg.write_text('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 400 200"><rect width="400" height="200" fill="white"/></svg>')
        from PIL import Image
        Image.new("RGB", (2000, 1000), "white").save(self.root / "figure.png")
        self.book = {"delivery_kind":"technical-smoke","title":"测试","reader":"读者","mode":"concept-analogy","scope":"软件测试夹具，不是读物", "sources":[
            {"id":f"S{i}","title":f"本地测试声明{i}","url":f"https://example.org/test/{i}","kind":"official" if i==1 else "reference", "verified":True,"retrieved_on":"2026-01-01","excerpt":"Fixture only: not production research evidence."} for i in range(1,4)],
            "claims":[{"id":"C01","text":"测试验证器行为","source_ids":["S1"],"status":"verified","critical":True}],
            "parts":[{"id":"P01","name":"夹具"}],
            "figures":[{"id":"F01","question":"测试什么","kind":"overview","caption":"测试图，不是正式科学插画","alt_text":"空白矩形夹具", "svg":"figure.svg","png":"figure.png","claim_ids":["C01"],"source_ids":["S1"],"part_ids":["P01"],"provenance":"original"}],
            "pages":[{"title":"验证器夹具","blocks":[{"type":"paragraph","text":"这段可编辑正文只用于本地软件测试，并不宣称外部资料已经回读。它包含足够字符供结构检查，不会作为用户样书交付。"*3},{"type":"figure","id":"F01"}]}]}
    def tearDown(self):
        self.temp.cleanup()
    def validate(self, book=None):
        path = self.root / "book.json"
        path.write_text(json.dumps(book or self.book,ensure_ascii=False),encoding="utf-8")
        return path, manual.validate_book(path)[1]
    def test_valid_structural_fixture(self):
        self.assertTrue(self.validate()[1]["ok"])
    def test_thin_fixture_cannot_be_production(self):
        b=copy.deepcopy(self.book); b["delivery_kind"]="production"
        self.assertFalse(self.validate(b)[1]["ok"])
    def test_missing_kind_defaults_to_production(self):
        b=copy.deepcopy(self.book); b.pop("delivery_kind")
        self.assertFalse(self.validate(b)[1]["ok"])
    def test_invalid_delivery_kind_rejected(self):
        b=copy.deepcopy(self.book); b["delivery_kind"]="pretend-production"
        self.assertFalse(self.validate(b)[1]["ok"])
    def test_complete_production_structure_accepted(self):
        b=copy.deepcopy(self.book); b["delivery_kind"]="production"
        kinds=["overview","exploded","section","detail","sequence","care"]
        b["figures"]=[]; b["pages"]=[]
        for i,kind in enumerate(kinds):
            f=copy.deepcopy(self.book["figures"][0]); f.update(id=f"F{i+1:02d}",kind=kind)
            b["figures"].append(f)
            p=copy.deepcopy(self.book["pages"][0]); p["blocks"][1]["id"]=f["id"]
            b["pages"].append(p)
        self.assertTrue(self.validate(b)[1]["ok"])
    def test_repeating_one_view_cannot_prove_production(self):
        b=copy.deepcopy(self.book); b["delivery_kind"]="production"
        b["pages"]=b["pages"]*6
        b["figures"]=[dict(b["figures"][0],id=f"F{i+1:02d}") for i in range(6)]
        b["pages"]=[copy.deepcopy(self.book["pages"][0]) for _ in range(6)]
        for i,p in enumerate(b["pages"]): p["blocks"][1]["id"]=f"F{i+1:02d}"
        self.assertFalse(self.validate(b)[1]["ok"])
    def test_unknown_source_rejected(self):
        b=copy.deepcopy(self.book); b["claims"][0]["source_ids"]=["S99"]
        self.assertFalse(self.validate(b)[1]["ok"])
    def test_unverified_critical_rejected(self):
        b=copy.deepcopy(self.book); b["claims"][0]["status"]="unverified"
        self.assertFalse(self.validate(b)[1]["ok"])
    def test_assumption_requires_notice_and_test(self):
        b=copy.deepcopy(self.book); b["claims"][0]["status"]="assumption"
        self.assertFalse(self.validate(b)[1]["ok"])
        b["claims"][0].update(test_evidence="test_model.py",assumption_notice="明确演示假设")
        self.assertTrue(self.validate(b)[1]["ok"])
    def test_path_escape_rejected(self):
        b=copy.deepcopy(self.book); b["figures"][0]["svg"]="../outside.svg"
        self.assertFalse(self.validate(b)[1]["ok"])
    def test_unknown_part_rejected(self):
        b=copy.deepcopy(self.book); b["figures"][0]["part_ids"]=["P404"]
        self.assertFalse(self.validate(b)[1]["ok"])
    def test_low_resolution_rejected(self):
        from PIL import Image
        Image.new("RGB",(200,100)).save(self.root/"figure.png")
        self.assertFalse(self.validate()[1]["ok"])
    def test_remote_svg_rejected(self):
        self.svg.write_text('<svg viewBox="0 0 4 2"><image href="https://example.org/a.png"/></svg>')
        with self.assertRaises(ValueError): manual.check_svg(self.svg)
    def test_svg_events_rejected(self):
        self.svg.write_text('<svg viewBox="0 0 4 2" onload="alert(1)"/>')
        with self.assertRaises(ValueError): manual.check_svg(self.svg)
    def test_svg_local_paint_reference_allowed(self):
        self.svg.write_text('<svg viewBox="0 0 4 2"><rect style="fill:url(#gradient)"/></svg>')
        manual.check_svg(self.svg)
    def test_docx_has_editable_body_and_media(self):
        path,_=self.validate(); out=self.root/"book.docx"
        manual.build(path,out)
        self.assertTrue(manual.inspect_docx(out)["ok"])
    def test_docx_metadata_is_neutral(self):
        path,_=self.validate(); out=self.root/"neutral.docx"
        manual.build(path,out)
        from docx import Document
        props=Document(out).core_properties
        for field in ("author","last_modified_by","comments","keywords"):
            self.assertEqual(getattr(props,field),"")
    def test_legacy_author_does_not_leak_into_docx(self):
        b=copy.deepcopy(self.book); b["author"]="Private Example Name"
        path,_=self.validate(b); out=self.root/"legacy.docx"
        manual.build(path,out)
        from docx import Document
        self.assertEqual(Document(out).core_properties.author,"")
    def test_qa_requires_render_and_review(self):
        path,_=self.validate(); out=self.root/"book.docx"
        manual.build(path,out)
        result=manual.qa_book(path,out,self.root/"qa")
        self.assertFalse(result["ok"])
    def test_source_floor_can_be_raised(self):
        b=copy.deepcopy(self.book);b['min_verified_sources']=4
        self.assertFalse(self.validate(b)[1]['ok'])
    def test_bad_source_floor_is_rejected_not_crash(self):
        b=copy.deepcopy(self.book);b['min_verified_sources']='many'
        self.assertFalse(self.validate(b)[1]['ok'])
    def test_v03_requires_asset_manifest(self):
        b=copy.deepcopy(self.book);b['schema_version']='0.3'
        self.assertFalse(self.validate(b)[1]['ok'])
    def test_sample_chapter_declares_its_limits(self):
        b=copy.deepcopy(self.book);b['delivery_kind']='sample-chapter'
        self.assertFalse(self.validate(b)[1]['ok'])
        b.update(sample_objectives=['fixture only'],sample_limits='not a complete book')
        self.assertTrue(self.validate(b)[1]['ok'])
    def test_context_image_does_not_need_fake_svg(self):
        b=copy.deepcopy(self.book);f=b['figures'][0]
        f.update(method='generated',role='context',claim_ids=[]);f.pop('svg')
        self.assertTrue(self.validate(b)[1]['ok'])
    def test_technical_svg_still_required(self):
        b=copy.deepcopy(self.book);b['figures'][0].pop('svg')
        self.assertFalse(self.validate(b)[1]['ok'])
    def test_raster_only_technical_keeps_claims_required(self):
        b=copy.deepcopy(self.book);f=b['figures'][0];f.update(method='existing',claim_ids=[]);f.pop('svg')
        self.assertFalse(self.validate(b)[1]['ok'])
    def test_fixed_break_attached_to_real_paragraph(self):
        b=copy.deepcopy(self.book);b.update(page_layout='fixed',include_sources=False)
        b['pages']=[copy.deepcopy(b['pages'][0]),copy.deepcopy(b['pages'][0])]
        b['pages'][1]['kicker']='fixture second page'
        path,_=self.validate(b);out=self.root/'break.docx';manual.build(path,out)
        from docx import Document
        d=Document(out);matches=[p for p in d.paragraphs if p.text=='fixture second page']
        self.assertTrue(matches[0].paragraph_format.page_break_before)
        self.assertFalse(any('w:type="page"' in p._p.xml for p in d.paragraphs))
    def test_template_cannot_pass_as_finished_book(self):
        _,result=manual.validate_book(HERE.parent/"templates/book.example.json")
        self.assertFalse(result["ok"])

if __name__ == "__main__":
    unittest.main(verbosity=2)
